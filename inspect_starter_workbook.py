import openpyxl

wb = openpyxl.load_workbook('Indian_IPO_ML_Research_Starter.xlsx')
for name in wb.sheetnames:
    sheet = wb[name]
    header = [cell.value for cell in sheet[1]] if sheet.max_row >= 1 else []
    print(f"Sheet '{name}': rows={sheet.max_row}, cols={sheet.max_column}")
    print("  Headers:", header[:12])
