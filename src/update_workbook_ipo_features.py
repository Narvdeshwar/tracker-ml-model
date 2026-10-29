import pandas as pd
import openpyxl

wb = openpyxl.load_workbook('Indian_IPO_ML_Research_Starter.xlsx')

sheet_name = 'IPO_Market_Features'
if sheet_name in wb.sheetnames:
    ws = wb[sheet_name]
    while ws.max_row > 0:
        ws.delete_rows(1)
else:
    ws = wb.create_sheet(title=sheet_name)

df_feat = pd.read_csv('data/processed/ipo_market_features.csv')

# Write headers
headers = df_feat.columns.tolist()
ws.append(headers)

# Write rows
for _, row in df_feat.iterrows():
    vals = []
    for h in headers:
        v = row[h]
        if pd.isna(v):
            vals.append(None)
        else:
            vals.append(v)
    ws.append(vals)

wb.save('Indian_IPO_ML_Research_Starter.xlsx')
print(f"Successfully saved {sheet_name} sheet with {ws.max_row - 1} records and {len(headers)} columns.")
