"""
Script: src/build_step_5c_datasets.py
Objective: Build standardized, timestamp-disciplined Subscription_Observations,
GMP_Observations, Experiment A (issue-close) and Experiment B (post-close) feature sets.
"""

import os
import pandas as pd
import numpy as np
import openpyxl

def build_step_5c_datasets():
    os.makedirs('data/processed', exist_ok=True)
    
    # 1. Load existing datasets
    df_market_feat = pd.read_csv('data/processed/ipo_market_features.csv')
    df_sent_feat = pd.read_csv('data/processed/ipo_sentiment_features.csv')
    
    # 2. Build Schema for Subscription_Observations and GMP_Observations
    # Since historical intra-day subscription and GMP logs with audited timestamps are not yet fully archived
    # in the verified repository, we instantiate the explicit observation tables with strict schema and provenance.
    # Any available observations must follow availability classification.
    
    sub_cols = [
        'ipo_id', 'observation_timestamp', 'observation_date', 'category',
        'subscription_multiple', 'source_name', 'source_url', 'retrieved_at',
        'availability_status', 'verification_status', 'notes'
    ]
    df_sub_obs = pd.DataFrame(columns=sub_cols)
    df_sub_obs.to_csv('data/processed/subscription_observations.csv', index=False)
    
    gmp_cols = [
        'ipo_id', 'observation_timestamp', 'observation_date', 'gmp_amount_inr',
        'issue_price_inr', 'gmp_pct_of_issue_price', 'source_name', 'source_url',
        'retrieved_at', 'availability_status', 'verification_status', 'notes'
    ]
    df_gmp_obs = pd.DataFrame(columns=gmp_cols)
    df_gmp_obs.to_csv('data/processed/gmp_observations.csv', index=False)
    
    # 3. Build Experiment A: Issue-Close Information Set
    # Cutoff: Exactly at Issue Close Date (end of bidding window, e.g., 17:00 IST)
    # Permitted features: Market conditions up to Issue Close date; Recent IPO sentiment up to Issue Close date.
    # Subscription / GMP: ONLY PRE_ISSUE_CLOSE or AT_ISSUE_CLOSE observations with verified timestamps.
    df_exp_a = df_market_feat.merge(
        df_sent_feat[[
            'ipo_id', 'recent5_mean_listing_open_return_pct', 'recent5_median_listing_open_return_pct',
            'recent5_positive_listing_rate', 'recent10_mean_listing_open_return_pct',
            'recent10_median_listing_open_return_pct', 'recent10_positive_listing_rate',
            'sentiment_missing_reason'
        ]],
        on='ipo_id',
        how='left'
    )
    df_exp_a['experiment_name'] = 'Experiment_A_Issue_Close'
    df_exp_a['prediction_timestamp_ist'] = df_exp_a['prediction_cutoff_date'].apply(
        lambda d: f"{d} 17:00:00" if pd.notnull(d) else None
    )
    # Subscription / GMP placeholders strictly audited
    df_exp_a['qib_subscription_multiple'] = np.nan
    df_exp_a['nii_subscription_multiple'] = np.nan
    df_exp_a['retail_subscription_multiple'] = np.nan
    df_exp_a['total_subscription_multiple'] = np.nan
    df_exp_a['subscription_availability_status'] = 'PENDING_HISTORICAL_TIMESTAMPS'
    df_exp_a['gmp_amount_inr'] = np.nan
    df_exp_a['gmp_pct_of_issue_price'] = np.nan
    df_exp_a['gmp_availability_status'] = 'PENDING_HISTORICAL_TIMESTAMPS'
    
    df_exp_a.to_csv('data/processed/ipo_features_issue_close.csv', index=False)
    print(f"Generated data/processed/ipo_features_issue_close.csv with {len(df_exp_a)} records.")
    
    # 4. Build Experiment B: Post-Close Information Set
    # Cutoff: After bidding closes and exchange basis of allotment is finalized (e.g., Listing-Eve 18:00 IST)
    df_exp_b = df_exp_a.copy()
    df_exp_b['experiment_name'] = 'Experiment_B_Post_Close'
    df_exp_b['prediction_timestamp_ist'] = df_exp_b['listing_date'].apply(
        lambda d: (pd.to_datetime(d) - pd.Timedelta(days=1)).strftime('%Y-%m-%d 18:00:00') if pd.notnull(d) else None
    )
    df_exp_b['subscription_availability_status'] = 'PENDING_FINAL_ALLOTMENT_VERIFICATION'
    df_exp_b['gmp_availability_status'] = 'PENDING_PRE_LISTING_GMP_AUDIT'
    
    df_exp_b.to_csv('data/processed/ipo_features_post_close.csv', index=False)
    print(f"Generated data/processed/ipo_features_post_close.csv with {len(df_exp_b)} records.")
    
    # 5. Update Workbook
    wb = openpyxl.load_workbook('Indian_IPO_ML_Research_Starter.xlsx')
    
    # Write ipo_sentiment_features
    sheet_name = 'IPO_Sentiment_Features'
    if sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        while ws.max_row > 0:
            ws.delete_rows(1)
    else:
        ws = wb.create_sheet(title=sheet_name)
    headers = df_sent_feat.columns.tolist()
    ws.append(headers)
    for _, r in df_sent_feat.iterrows():
        ws.append([None if pd.isna(r[h]) else r[h] for h in headers])
        
    wb.save('Indian_IPO_ML_Research_Starter.xlsx')
    print("Workbook updated with IPO_Sentiment_Features sheet.")

if __name__ == '__main__':
    build_step_5c_datasets()
