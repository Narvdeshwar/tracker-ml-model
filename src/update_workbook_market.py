import pandas as pd
import openpyxl

wb = openpyxl.load_workbook('Indian_IPO_ML_Research_Starter.xlsx')
ws_market = wb['Market_Daily']
headers = [c.value for c in ws_market[1]]
print('Existing Market_Daily headers in Excel:', headers)

df_m = pd.read_csv('data/processed/market_daily.csv')

# Clear existing rows beyond header
while ws_market.max_row > 1:
    ws_market.delete_rows(2)

# Map our columns to workbook columns
# Headers: ['date', 'nifty50_close', 'nifty50_return_1d_pct', 'nifty50_return_5d_pct', 'nifty50_return_20d_pct', 'nifty500_return_5d_pct', 'nifty500_return_20d_pct', 'india_vix_close', 'india_vix_change_5d_pct', 'advance_decline_ratio', 'fii_net_cash_cr', 'dii_net_cash_cr']

for _, r in df_m.iterrows():
    row_data = {
        'date': r['date'],
        'nifty50_close': r['nifty50_close'],
        'nifty50_return_1d_pct': round(r['nifty50_return_1d'] * 100, 4) if pd.notnull(r['nifty50_return_1d']) else None,
        'nifty50_return_5d_pct': round(r['nifty50_return_5s'] * 100, 4) if pd.notnull(r['nifty50_return_5s']) else None,
        'nifty50_return_20d_pct': round(r['nifty50_return_20s'] * 100, 4) if pd.notnull(r['nifty50_return_20s']) else None,
        'nifty500_return_5d_pct': round(r['nifty500_return_5s'] * 100, 4) if pd.notnull(r['nifty500_return_5s']) else None,
        'nifty500_return_20d_pct': round(r['nifty500_return_20s'] * 100, 4) if pd.notnull(r['nifty500_return_20s']) else None,
        'india_vix_close': r['india_vix_close'],
        'india_vix_change_5d_pct': round(r['india_vix_pct_change_5s'] * 100, 4) if pd.notnull(r['india_vix_pct_change_5s']) else None,
        'advance_decline_ratio': None,
        'fii_net_cash_cr': None,
        'dii_net_cash_cr': None
    }
    ws_market.append([row_data.get(col, '') for col in headers])

wb.save('Indian_IPO_ML_Research_Starter.xlsx')
print(f"Successfully updated Market_Daily sheet with {ws_market.max_row - 1} trading session records.")
